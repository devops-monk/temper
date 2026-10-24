"""Data Recipe worker: `python recipe_worker.py <job.json>`.

uv runs this file by path in its own environment with NVIDIA's data-designer (Apache-2.0) and the
document parsers, so the engine itself stays small. It must not import `temper_engine`: none of the
engine's packages exist in that environment. Prints `@@{json}` lines like the other workers.
"""

from __future__ import annotations

import io
import json
import logging
import re
import sys
import urllib.request
import zipfile
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".html", ".htm", ".xml"}
CODE_SUFFIXES = {".py", ".js", ".ts", ".tsx", ".jsx", ".go", ".rs", ".java", ".kt", ".swift", ".c", ".h", ".cpp", ".hpp", ".cs", ".rb", ".php", ".sh", ".sql", ".toml"}
DOC_SUFFIXES = TEXT_SUFFIXES | CODE_SUFFIXES | {".pdf", ".docx"}
MAX_FILE_BYTES = 2_000_000  # skip generated or vendored blobs in repositories
GITHUB_REPO = re.compile(r"^https?://github\.com/([\w.-]+)/([\w.-]+?)(?:\.git)?/?$")
USER_AGENT = "Temper data recipes"


def emit(event: str, **data) -> None:
    print("@@" + json.dumps({"event": event, **data}, default=str), flush=True)


# ------------------------------------------------------------------ documents → text


def html_text(raw: bytes) -> str:
    import lxml.html

    tree = lxml.html.fromstring(raw)
    for node in tree.xpath("//script|//style|//noscript|//nav|//footer|//header"):
        node.drop_tree()
    lines = (line.strip() for line in tree.text_content().splitlines())
    return "\n".join(line for line in lines if line)


def file_text(name: str, data: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages)
    if suffix == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
        return "\n".join(p for p in parts if p.strip())
    if suffix in (".html", ".htm"):
        return html_text(data)
    return data.decode("utf-8", errors="replace")


def fetch(url: str) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=60) as response:
        return response.read()


def github_repo(owner: str, repo: str) -> list[tuple[str, str]]:
    """Text and code files from the default branch of a public repository."""
    archive = zipfile.ZipFile(io.BytesIO(fetch(f"https://codeload.github.com/{owner}/{repo}/zip/HEAD")))
    out = []
    for info in archive.infolist():
        name = info.filename.split("/", 1)[-1]  # drop the "<repo>-<sha>/" prefix
        if info.is_dir() or info.file_size > MAX_FILE_BYTES or Path(name).suffix.lower() not in TEXT_SUFFIXES | CODE_SUFFIXES:
            continue
        if any(part in {"node_modules", "vendor", "dist", "build", ".git"} for part in Path(name).parts):
            continue
        out.append((f"{owner}/{repo}/{name}", archive.read(info).decode("utf-8", errors="replace")))
    return out


def load_documents(sources: list[str]) -> list[tuple[str, str]]:
    """(source, text) for every file, folder, web page or GitHub repository in `sources`."""
    docs: list[tuple[str, str]] = []
    for source in sources:
        source = source.strip()
        if not source:
            continue
        print(f"Reading {source}", flush=True)
        if m := GITHUB_REPO.match(source):
            docs += github_repo(m.group(1), m.group(2))
        elif source.startswith(("http://", "https://")):
            data = fetch(source)
            path = source.split("?")[0]
            docs.append((source, file_text(path, data) if Path(path).suffix.lower() in {".pdf", ".docx"} else html_text(data)))
        else:
            path = Path(source).expanduser()
            files = sorted(p for p in path.rglob("*") if p.is_file()) if path.is_dir() else [path]
            for file in files:
                if file.suffix.lower() in DOC_SUFFIXES and file.stat().st_size <= MAX_FILE_BYTES * 25:
                    docs.append((str(file), file_text(file.name, file.read_bytes())))
    return [(name, text) for name, text in docs if text.strip()]


def chunk(text: str, size: int, overlap: int) -> list[str]:
    """Split into chunks of about `size` characters, breaking at paragraph, line or word boundaries."""
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            window = text[start:end]
            cut = max(window.rfind("\n\n"), window.rfind("\n"), window.rfind(". "), window.rfind(" "))
            if cut > size // 2:
                end = start + cut + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in chunks if c]


def document_seed(seed: dict, work: Path) -> Path:
    docs = load_documents(seed.get("sources", []))
    if not docs:
        raise RuntimeError("None of the sources had any text. Check the paths and links.")
    size, overlap = int(seed.get("chunk_size", 1200)), int(seed.get("chunk_overlap", 200))
    rows = [{"chunk_text": c, "source_file": name} for name, text in docs for c in chunk(text, size, overlap)]
    path = work / "document_chunks.jsonl"
    with open(path, "w", encoding="utf-8") as f:
        f.writelines(json.dumps(r, ensure_ascii=False) + "\n" for r in rows)
    emit("seed", documents=len(docs), chunks=len(rows), sample=rows[0]["chunk_text"][:600])
    return path


# ------------------------------------------------------------------ recipe → data-designer config


def build_config(recipe: dict, work: Path) -> dict:
    models = [
        {
            "alias": m["alias"],
            "model": m.get("model") or "local",
            "provider": m.get("provider") or "local",
            "skip_health_check": True,
            "inference_parameters": {
                k: m[k] for k in ("temperature", "top_p", "max_tokens", "max_parallel_requests", "timeout") if m.get(k) is not None
            },
        }
        for m in recipe.get("models", [])
    ]
    config: dict = {"columns": recipe["columns"], "model_configs": models, "processors": recipe.get("processors", [])}
    seed = recipe.get("seed")
    if seed:
        kind = seed.get("type")
        if kind == "documents":
            source = {"seed_type": "local", "path": str(document_seed(seed, work))}
        elif kind == "file":
            source = {"seed_type": "local", "path": str(Path(seed["path"]).expanduser())}
        elif kind == "hf":
            path = seed["path"]
            # "owner/name" means every parquet file of that dataset.
            source = {"seed_type": "hf", "path": path if path.startswith("datasets/") else f"datasets/{path}/**/*.parquet"}
        else:
            raise RuntimeError(f"Unknown seed type: {kind}")
        config["seed_config"] = {"source": source, "sampling_strategy": seed.get("sampling", "ordered")}
    return config


def providers(job: dict) -> list:
    from data_designer.config import ModelProvider

    out = []
    if job.get("local_endpoint"):
        out.append(ModelProvider(name="local", endpoint=job["local_endpoint"], provider_type="openai", api_key="local"))
    for p in job["recipe"].get("providers", []):
        out.append(ModelProvider(name=p["name"], endpoint=p["endpoint"], provider_type="openai", api_key=p.get("api_key") or None))
    return out


# ------------------------------------------------------------------ progress


class Progress(logging.Handler):
    """Turns data-designer's progress log lines into `progress` events.

    It logs "(2/5) Dispatching ..." per row group and "column 'x': 3/10 (30%)" per column within it.
    """

    GROUP = re.compile(r"\((\d+)/(\d+)\) Dispatching")
    COLUMN = re.compile(r"column '([^']+)': (\d+)/(\d+)")

    def __init__(self) -> None:
        super().__init__()
        self.group, self.groups, self.columns = 1, 1, {}

    def emit(self, record: logging.LogRecord) -> None:
        message = record.getMessage()
        if m := self.GROUP.search(message):
            self.group, self.groups, self.columns = int(m.group(1)), int(m.group(2)), {}
        elif m := self.COLUMN.search(message):
            self.columns[m.group(1)] = (int(m.group(2)), int(m.group(3)))
            done = sum(d for d, _ in self.columns.values()) / max(sum(t for _, t in self.columns.values()), 1)
            emit("progress", value=min((self.group - 1 + done) / self.groups, 0.99))


def records(frame) -> list[dict]:
    return json.loads(frame.to_json(orient="records", force_ascii=False, default_handler=str))


def main() -> None:
    job_path = Path(sys.argv[1])
    job = json.loads(job_path.read_text())
    work = job_path.parent
    try:
        from data_designer.config import DataDesignerConfigBuilder
        from data_designer.interface.data_designer import DataDesigner

        emit("phase", name="Reading sources" if (job["recipe"].get("seed") or {}).get("type") == "documents" else "Preparing")
        builder = DataDesignerConfigBuilder.from_config(build_config(job["recipe"], work))
        designer = DataDesigner(work / "artifacts", model_providers=providers(job))
        logging.getLogger().addHandler(Progress())
        emit("phase", name="Generating")
        if job["mode"] == "preview":
            frame = designer.preview(builder, num_records=job["num_records"]).dataset
        else:
            frame = designer.create(builder, num_records=job["num_records"], dataset_name="recipe").load_dataset()
        rows = records(frame)
        if job.get("output"):
            out = Path(job["output"])
            out.parent.mkdir(parents=True, exist_ok=True)
            with open(out, "w", encoding="utf-8") as f:
                f.writelines(json.dumps(r, ensure_ascii=False) + "\n" for r in rows)
        emit("result", rows=rows[:50], count=len(rows), columns=list(frame.columns), output=job.get("output"))
    except Exception as exc:
        emit("error", message=f"{type(exc).__name__}: {exc}")
        raise


if __name__ == "__main__":
    main()
